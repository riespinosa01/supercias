"""Transform tests that do not call Supercias."""

from __future__ import annotations

import io
import logging
import tempfile
import unittest
import zipfile
from pathlib import Path
from unittest import mock

import pandas as pd
from openpyxl import Workbook

from supercias.cli import column_catalog, main, parse_columns
from supercias.columns import DEFAULT_COLUMNS
from supercias.directorio import read_directorio_xlsx
from supercias.ranking import ranking_frame_from_text
from supercias.transform import build_dataset, select_balance_year


def _directory_rows() -> pd.DataFrame:
    rows = []
    for index in range(100):
        rows.append(
            {
                "EXPEDIENTE": index + 1,
                "RUC": f"179000000{index:04d}",
                "NOMBRE": f"COMPANIA {index + 1}",
                "SITUACIÓN LEGAL": "ACTIVA",
                "FECHA_CONSTITUCION": "01/01/2010",
                "TIPO": "ANÓNIMA",
                "PAÍS": "ECUADOR",
                "PROVINCIA": "PICHINCHA",
                "CIUDAD": "QUITO",
                "ÚLTIMO BALANCE": 2025,
                "CIIU NIVEL 1": "G",
                "CIIU NIVEL 6": "G4610.03",
                "PRESENTÓ BALANCE INICIAL": "NO APLICA" if index % 2 == 0 else "SI",
            }
        )
    rows.append(
        {
            "EXPEDIENTE": 9000,
            "RUC": "0990000000001",
            "NOMBRE": "INICIAL 2026",
            "SITUACIÓN LEGAL": "ACTIVA",
            "FECHA_CONSTITUCION": "12/02/2026",
            "TIPO": "SOCIEDAD POR ACCIONES SIMPLIFICADA",
            "PAÍS": "ECUADOR",
            "PROVINCIA": "GUAYAS",
            "CIUDAD": "GUAYAQUIL",
            "ÚLTIMO BALANCE": 2026,
            "CIIU NIVEL 1": "L",
            "CIIU NIVEL 6": "L6810.01",
            "PRESENTÓ BALANCE INICIAL": "SI",
        }
    )
    rows.append(
        {
            "EXPEDIENTE": 8000,
            "RUC": "0990000000002",
            "NOMBRE": "DISUELTA",
            "SITUACIÓN LEGAL": "INACTIVA",
            "FECHA_CONSTITUCION": "01/01/2000",
            "TIPO": "ANÓNIMA",
            "PAÍS": "ECUADOR",
            "PROVINCIA": "PICHINCHA",
            "CIUDAD": "QUITO",
            "ÚLTIMO BALANCE": 2025,
            "CIIU NIVEL 1": "G",
            "CIIU NIVEL 6": "G4610.03",
            "PRESENTÓ BALANCE INICIAL": "NO",
        }
    )
    return pd.DataFrame(rows)


def _ranking_rows() -> pd.DataFrame:
    rows = []
    for index in range(100):
        expediente = index + 1
        # Mean utilidad across the two years: (40_000 + 60_000) / 2 = 50_000
        # for even companies, and (4_000_000 + 6_000_000) / 2 = 5_000_000.
        # Company 3 (index 2) is below the floor. Company 4 (index 3) is above the ceiling.
        if index == 2:
            utilities = (10_000, 20_000)
        elif index == 3:
            utilities = (6_000_000, 8_000_000)
        elif index % 2 == 0:
            utilities = (40_000, 60_000)
        else:
            utilities = (4_000_000, 6_000_000)
        for year, utility in ((2024, utilities[0]), (2025, utilities[1])):
            rows.append(
                {
                    "anio": year,
                    "expediente": expediente,
                    "ingresos_totales": 200_000 if utility < 100_000 else 8_000_000,
                    "costos_ventas_prod": 50_000,
                    "gastos_admin_ventas": 10_000,
                    "gastos_financieros": 1_000,
                    "utilidad_neta": utility,
                    "deuda_total": 5_000,
                    "deuda_total_c_plazo": 1_000,
                    "n_empleados": 4,
                }
            )
    # Incomplete financials must be dropped.
    rows.append(
        {
            "anio": 2025,
            "expediente": 1,
            "ingresos_totales": 10,
            "costos_ventas_prod": 1,
            "gastos_admin_ventas": 1,
            "gastos_financieros": 1,
            "utilidad_neta": None,
            "deuda_total": 1,
            "deuda_total_c_plazo": 1,
            "n_empleados": 1,
        }
    )
    return pd.DataFrame(rows)


class TransformTests(unittest.TestCase):
    def test_latest_cohort_ignores_tiny_future_tail(self) -> None:
        years = pd.Series([2025] * 100 + [2026])
        choice = select_balance_year(years)
        self.assertEqual(choice.literal_max, 2026)
        self.assertEqual(choice.selected, 2025)
        self.assertEqual(choice.literal_rows, 1)
        self.assertEqual(choice.selected_rows, 100)

    def test_balance_year_override(self) -> None:
        years = pd.Series([2025] * 100 + [2026])
        choice = select_balance_year(years, override=2026)
        self.assertEqual(choice.selected, 2026)

    def test_default_export_filters_and_adds_margen(self) -> None:
        frame = build_dataset(_directory_rows(), _ranking_rows())
        self.assertEqual(list(frame.columns), [*DEFAULT_COLUMNS, "margen_neto"])
        self.assertNotIn(8000, set(frame["EXPEDIENTE"]))
        self.assertNotIn(9000, set(frame["EXPEDIENTE"]))
        self.assertNotIn(3, set(frame["EXPEDIENTE"]))
        self.assertNotIn(4, set(frame["EXPEDIENTE"]))
        self.assertIn(1, set(frame["EXPEDIENTE"]))
        company = frame.loc[frame["EXPEDIENTE"] == 1]
        self.assertEqual(len(company), 2)
        self.assertEqual(int(company["empresa_nueva"].iloc[0]), 0)
        odd = frame.loc[frame["EXPEDIENTE"] == 2]
        self.assertEqual(int(odd["empresa_nueva"].iloc[0]), 1)
        self.assertAlmostEqual(float(company["margen_neto"].iloc[0]), 40_000 / 200_000)

    def test_bounds_are_inclusive(self) -> None:
        frame = build_dataset(_directory_rows(), _ranking_rows(), floor=50_000, ceiling=5_000_000)
        means = frame.groupby("RUC")["utilidad_neta"].mean()
        self.assertGreaterEqual(means.min(), 50_000)
        self.assertLessEqual(means.max(), 5_000_000)

    def test_skip_filter_keeps_out_of_band_companies(self) -> None:
        frame = build_dataset(_directory_rows(), _ranking_rows(), skip_filter=True)
        self.assertIn(3, set(frame["EXPEDIENTE"]))
        self.assertIn(4, set(frame["EXPEDIENTE"]))
        self.assertNotIn(9000, set(frame["EXPEDIENTE"]))

    def test_column_override_and_zero_revenue_margen(self) -> None:
        directory = _directory_rows().head(1).copy()
        ranking = pd.DataFrame(
            [
                {
                    "anio": 2025,
                    "expediente": 1,
                    "ingresos_totales": 0,
                    "utilidad_neta": 10,
                    "deuda_total": 1,
                    "deuda_total_c_plazo": 1,
                    "n_empleados": 2,
                    "costos_ventas_prod": 0,
                    "gastos_admin_ventas": 0,
                    "gastos_financieros": 0,
                }
            ]
        )
        frame = build_dataset(
            directory,
            ranking,
            columns=["EXPEDIENTE", "NOMBRE", "ingresos_totales", "utilidad_neta"],
            floor=0,
            ceiling=100,
        )
        self.assertEqual(
            list(frame.columns),
            ["EXPEDIENTE", "NOMBRE", "ingresos_totales", "utilidad_neta", "margen_neto"],
        )
        self.assertTrue(pd.isna(frame["margen_neto"].iloc[0]))

    def test_published_empresa_nueva_is_kept(self) -> None:
        directory = _directory_rows().head(1).copy()
        directory["PRESENTÓ BALANCE INICIAL"] = "SI"
        ranking = _ranking_rows()
        ranking = ranking.loc[ranking["expediente"] == 1].dropna()
        ranking["empresa_nueva"] = 0
        frame = build_dataset(directory, ranking, floor=0, ceiling=10_000_000)
        self.assertEqual(int(frame["empresa_nueva"].iloc[0]), 0)

    def test_accented_column_names_are_matched(self) -> None:
        directory = _directory_rows().rename(
            columns={"PAÍS": "PAIS", "ÚLTIMO BALANCE": "ULTIMO BALANCE", "SITUACIÓN LEGAL": "SITUACION LEGAL"}
        )
        frame = build_dataset(directory, _ranking_rows(), columns=["EXPEDIENTE", "PAÍS", "ÚLTIMO BALANCE"])
        self.assertEqual(list(frame.columns), ["EXPEDIENTE", "PAÍS", "ÚLTIMO BALANCE"])
        self.assertEqual(int(frame["ÚLTIMO BALANCE"].iloc[0]), 2025)

    def test_ranking_text_is_parsed_from_the_response_body(self) -> None:
        body = (
            "anio,expediente,ingresos_totales,utilidad_neta\n"
            "2025,10,100.50,20.25\n"
        )
        frame = ranking_frame_from_text(
            body,
            needed={"anio", "expediente", "ingresos_totales"},
            expedientes={10},
        )
        self.assertEqual(list(frame.columns), ["anio", "expediente", "ingresos_totales"])
        self.assertEqual(int(frame["expediente"].iloc[0]), 10)
        skipped = ranking_frame_from_text(body, expedientes={99})
        self.assertTrue(skipped.empty)
        semicolon = "anio;expediente;ingresos_totales\n2024;3;8\n"
        parsed = ranking_frame_from_text(semicolon)
        self.assertEqual(int(parsed["expediente"].iloc[0]), 3)

    def test_broken_worksheet_dimension_still_reads_every_column(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            path = Path(temporary) / "directorio.xlsx"
            _write_directory_workbook(path)
            _collapse_dimension(path)
            frame = read_directorio_xlsx(path)
        self.assertEqual(
            list(frame.columns[:4]),
            ["No. FILA", "EXPEDIENTE", "RUC", "NOMBRE"],
        )
        self.assertEqual(int(frame["EXPEDIENTE"].iloc[0]), 10)
        self.assertEqual(frame["SITUACIÓN LEGAL"].iloc[0], "ACTIVA")
        self.assertEqual(str(frame["RUC"].iloc[0]), "1790013731001")

    def test_help_renders(self) -> None:
        with self.assertRaises(SystemExit) as caught:
            main(["--help"])
        self.assertEqual(caught.exception.code, 0)

    def test_list_columns_and_parser(self) -> None:
        buffer = io.StringIO()
        root = logging.getLogger()
        previous_handlers = root.handlers[:]
        previous_level = root.level
        try:
            with mock.patch("sys.stdout", buffer):
                status = main(["--list-columns"])
        finally:
            root.handlers[:] = previous_handlers
            root.setLevel(previous_level)
        self.assertEqual(status, 0)
        text = buffer.getvalue()
        self.assertIn("EXPEDIENTE", text)
        self.assertIn("utilidad_neta", text)
        self.assertIn("margen_neto", text)
        self.assertIn("empresa_nueva", text)
        self.assertEqual(
            parse_columns(["EXPEDIENTE, RUC", "NOMBRE", "RUC"]),
            ["EXPEDIENTE", "RUC", "NOMBRE"],
        )
        self.assertIn("PAÍS", column_catalog())

    def test_dropna_removes_blank_expediente(self) -> None:
        directory = _directory_rows().head(1).copy()
        ranking = _ranking_rows()
        ranking.loc[ranking["expediente"] == 1, "expediente"] = None
        extra = ranking.iloc[0].copy()
        extra["expediente"] = None
        frame = build_dataset(directory, pd.concat([ranking, extra.to_frame().T], ignore_index=True), floor=0, ceiling=10_000_000)
        self.assertTrue(frame.empty)


def _write_directory_workbook(path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet["A1"] = "SUPERINTENDENCIA DE COMPAÑÍAS, VALORES Y SEGUROS"
    sheet["A2"] = "DIRECTORIO DE COMPAÑÍAS"
    headers = [
        "No. FILA",
        "EXPEDIENTE",
        "RUC",
        "NOMBRE",
        "SITUACIÓN LEGAL",
        "ÚLTIMO BALANCE",
    ]
    for column, header in enumerate(headers, start=1):
        sheet.cell(3, column, header)
    values = [1, 10, "1790013731001", "ACEITES", "ACTIVA", 2025]
    for column, value in enumerate(values, start=1):
        sheet.cell(4, column, value)
    workbook.save(path)


def _collapse_dimension(path: Path) -> None:
    original = path.read_bytes()
    path.unlink()
    with zipfile.ZipFile(io.BytesIO(original)) as source, zipfile.ZipFile(path, "w") as target:
        for info in source.infolist():
            payload = source.read(info.filename)
            if info.filename.startswith("xl/worksheets/sheet"):
                payload = payload.replace(b"<dimension ", b"<dimension ", 1)
                start = payload.find(b"<dimension ")
                end = payload.find(b"/>", start)
                payload = payload[:start] + b'<dimension ref="A1"/>' + payload[end + 2 :]
            target.writestr(info, payload)


if __name__ == "__main__":
    unittest.main()
